#!/usr/bin/env python3
"""Build CPI_NQ_Backtest_Template.xlsx.

Import it into the existing Google Sheet with File > Import > Upload >
"Insert new sheet(s)". Formulas only use functions that behave the same in
Google Sheets, Excel and LibreOffice (no array constants, no wildcards).
"""
from __future__ import annotations

import argparse
import datetime as dt
import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

LOG = "CPI Log"
STATS = "Stats"
GUIDE = "Guide"
FIRST_ROW = 4  # rows 1-3 = section band, headers, hints
LAST_ROW = 503  # 500 ready-made rows
STATS_LAST_ROW = 1000  # Stats reads further in case rows get added

# ---------------------------------------------------------------- options
# Labels are used as COUNTIFS criteria: never use * ? ~ , or a leading < > =.
WICK = ["Up", "Down", "No Wick"]
MAIN = ["Up", "Down"]
STRADDLE = ["TP", "SL", "BE", "N/A"]
CPI = ["Hotter", "Cooler", "In line", "Mixed"]
CLOSE = ["Top", "Middle", "Bottom"]
WIN_OPEN = ["0-15 sec", "15-30 sec", "30-60 sec", "1-5 min", "5-15 min",
            "15-60 min", "9:30-11:00", "11:00-16:00", "Not by 16:00"]
WIN_HL = WIN_OPEN[3:]  # the 1-min candle's high/low can only be taken after 8:31
NEVER = "Not by 16:00"
AT_OPEN = ["Swept & resumed", "Broke through", "Chopped around", "Never came back"]
BIAS = ["Bullish", "Bearish", "No bias"]
MIDNIGHT = ["Above (premium)", "Below (discount)"]
LIQ = ["None", "Pre-news high/low", "Asia high/low", "London high/low",
       "Prev day high/low", "Prev week high/low"]
PDA = ["No", "FVG", "Order block", "Breaker", "Other"]
FVG = ["No FVG", "Never came back", "Came back & held", "Came back & failed (IFVG)"]
SMT = ["Yes", "No", "Not checked"]
SETUP = ["None", "A Continuation", "B Reversal", "C Open sweep", "D 9:30-11:00 raid"]

OUTCOMES = ["Continued", "Back to open", "Partial retrace", "Stalled / chop"]
MOVE_VS_DATA = ["With data", "Against data", "Neutral data"]
HIGH_MAIN, LOW_MAIN = "High first (main side)", "Low first (main side)"
HIGH_WICK, LOW_WICK = "High first (wick side)", "Low first (wick side)"
SAME_WINDOW, NEITHER = "Same window", "Neither taken"

# ---------------------------------------------------------------- styling
FONT = "Arial"
GREEN, RED, AMBER, GREY = "BBF7D0", "F4CCCC", "FEF3C7", "E5E7EB"
AUTO_FILL = "F3F4F6"
SECTIONS = {
    # key: (band text, band colour, header colour, header font colour)
    "legacy": ("1 · YOUR COLUMNS", "111827", "374151", "FFFFFF"),
    "cpi": ("2 · CPI NUMBER", "1D4ED8", "DBEAFE", "111827"),
    "close": ("3 · CANDLE CLOSE", "6D28D9", "EDE9FE", "111827"),
    "after": ("4 · AFTER THE 1st MINUTE", "047857", "D1FAE5", "111827"),
    "ict": ("5 · ICT CONTEXT", "B91C1C", "FEE2E2", "111827"),
    "trade": ("6 · TRADE & REVIEW", "B45309", "FEF3C7", "111827"),
}
# Band text for the part of section 1 right of the frozen Date/Day columns.
LEGACY_BAND_REST = "same order as your old tab → paste your old rows at A4 with Ctrl+Shift+V (values only)"


def fill(hex_colour: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=hex_colour, end_color=hex_colour)


THIN = Side(style="thin", color="D1D5DB")
SECTION_EDGE = Side(style="medium", color="6B7280")


@dataclass
class Col:
    key: str
    header: str
    hint: str
    width: float
    section: str
    meaning: str
    how: str
    example: object = None
    options: list | None = None
    auto: bool = False
    number_format: str | None = None
    validation: str | None = None  # "pullback" | "r"
    letter: str = ""


COLUMNS = [
    # 1 · existing columns, same order as the original tab so one paste migrates it
    Col("date", "CPI Release Date", "date", 16, "legacy",
        "Day the CPI number came out (8:30 New York time).",
        "Type the date, e.g. 12 Aug 2026.", "(release date)", number_format="dd mmm yyyy"),
    Col("day", "Day", "auto", 7, "legacy",
        "Weekday of the release.",
        "Nothing - fills itself. (Old rows you paste keep the day you typed.)", "Wed", auto=True),
    Col("wick", "Wick", "▼ Up / Down / No Wick", 11, "legacy",
        "Direction of the first spike when the number hits (the quick move AGAINST the main move).",
        "▼ Up / Down, or No Wick if price went straight in the main direction.", "Down", WICK),
    Col("wick_pts", "Wick pts", "pts", 9, "legacy",
        "Size of that first spike in points.",
        "Number. Leave empty if No Wick.", 12),
    Col("main", "Main Move", "▼ Up / Down", 11, "legacy",
        "Direction price really went after the wick.",
        "▼ Up / Down. Most other columns are read relative to this.", "Up", MAIN),
    Col("m15", "15sec move", "pts", 10, "legacy",
        "How far the main move went in the first 15 seconds (points) - keep your usual method.",
        "Number only. If nothing changed, type the same number again instead of 'Same' (text can't be averaged).", 80),
    Col("m30", "30sec move", "pts", 10, "legacy",
        "Same, for the first 30 seconds.",
        "Number only (no 'Same'). A move the other way is captured by 'Deepest pullback %'.", 95),
    Col("m1", "1min move", "pts", 10, "legacy",
        "Same, for the first minute.",
        "Number only (no 'Same').", 110),
    Col("straddle", "Straddle", "▼ TP / SL / BE / N/A", 10, "legacy",
        "Result of your straddle order.",
        "▼ TP / SL / BE / N/A.", "TP", STRADDLE),
    Col("scam", "Scam Wick (15 sec before)", "pts", 13, "legacy",
        "Size of the fake move in the 15 seconds before the release (points).",
        "Number.", 6),
    Col("notes", "Notes", "text", 30, "legacy",
        "Anything unusual: delayed release, FOMC week, other news at 8:30, crazy spread...",
        "Free text.", "EXAMPLE ONLY"),
    # 2 · CPI number
    Col("cpi", "CPI vs forecast", "▼ Hotter / Cooler / In line / Mixed", 13, "cpi",
        "Was CPI hotter or cooler than expected? Compare CPI m/m and Core CPI m/m, actual vs forecast (ForexFactory history).",
        "▼ Hotter = both higher, or one higher + one equal · Cooler = both lower, or one lower + one equal · "
        "In line = both equal · Mixed = one higher, one lower.", "Cooler", CPI),
    Col("move_vs_data", "Move vs data (auto)", "auto", 13, "cpi",
        "Did NQ move the 'logical' way? Hot CPI usually = NQ down, cool CPI = NQ up.",
        "Nothing - auto: With data / Against data / Neutral data (in line or mixed).", "With data", auto=True),
    # 3 · candle close location
    Col("close1", "1m candle close", "▼ Top / Middle / Bottom", 11, "close",
        "Where the 8:30 one-minute candle closed inside its own range (wicks included).",
        "▼ Split the candle high-to-low into 3 equal parts: Top / Middle / Bottom third.", "Top", CLOSE),
    Col("close5", "5m candle close", "▼ Top / Middle / Bottom", 11, "close",
        "Where the 5-minute candle 8:30-8:34 closed inside its range.",
        "▼ Top / Middle / Bottom third.", "Top", CLOSE),
    Col("close15", "15m candle close", "▼ Top / Middle / Bottom", 11, "close",
        "Where the 15-minute candle 8:30-8:44 closed inside its range.",
        "▼ Top / Middle / Bottom third.", "Middle", CLOSE),
    # 4 · after the first minute
    Col("high_when", "1-min HIGH taken (BSL) – when?", "▼ time window", 15, "after",
        "When price first traded above the HIGH of the 8:30 1-min candle (buy-side liquidity).",
        "▼ Time window in NY time. 'Not by 16:00' if it never happened that day.", "5-15 min", WIN_HL),
    Col("low_when", "1-min LOW taken (SSL) – when?", "▼ time window", 15, "after",
        "When price first traded below the LOW of the 8:30 1-min candle (sell-side liquidity).",
        "▼ Time window in NY time. 'Not by 16:00' if it never happened that day.", NEVER, WIN_HL),
    Col("pullback", "Deepest pullback %", "0-200 (no % sign)", 12, "after",
        "How far price came back from the 1-min main-side extreme toward the 8:30 open, BEFORE it broke that extreme "
        "(stop checking at 8:45). 0 = no pullback · 50 = half-way · 62-79 = OTE · 100 = back at the 8:30 open · "
        "120 = 20% past the open.",
        "Number only: 62, not 62%. Fib with 0 on the 1-min main-side extreme and 100 on the 8:30 open.", 62,
        number_format='0"%"', validation="pullback"),
    Col("back_open", "Back to 8:30 open – when?", "▼ time window", 15, "after",
        "First time price traded back at the 8:30 open after the main move started.",
        "▼ Time window in NY time. 'Not by 16:00' if it never came back.", "15-60 min", WIN_OPEN),
    Col("at_open", "At the open, price…", "▼ pick", 16, "after",
        "What price did once it came back to the 8:30 open - does it just grab the liquidity there and keep going?",
        "▼ Swept & resumed = turned and made a new high/low in the main direction before taking the 1-min wick side · "
        "Broke through = took the 1-min wick side first · Chopped around = neither by 16:00 · Never came back.",
        "Swept & resumed", AT_OPEN),
    Col("outcome", "Outcome 8:31–8:45 (auto)", "auto", 16, "after",
        "What the main move did after the first minute.",
        "Nothing - auto from Main Move, Deepest pullback % and the main-side 'taken – when': "
        "Back to open = pullback reached 100 · Continued = main side taken by 8:45 (1-5 or 5-15 min) · "
        "Partial retrace = not taken by 8:45, pullback 50-99 · Stalled / chop = not taken by 8:45, pullback under 50.",
        "Continued", auto=True),
    Col("first_side", "First side taken (auto)", "auto", 18, "after",
        "Which end of the 1-min candle got taken first, and whether that was the main-move side or the wick side.",
        "Nothing - auto from the two 'taken – when' columns.", HIGH_MAIN, auto=True),
    # 5 · ICT context
    Col("bias", "Daily bias (before 8:30)", "▼ pick", 13, "ict",
        "Your higher-timeframe bias for the day, decided BEFORE you replay the release.",
        "▼ Bullish / Bearish / No bias.", "Bullish", BIAS),
    Col("midnight", "8:30 open vs Midnight open", "▼ Above / Below", 15, "ict",
        "Is the 8:30 open above the 00:00 New York opening price (premium) or below it (discount)?",
        "▼ Above (premium) / Below (discount).", "Below (discount)", MIDNIGHT),
    Col("liq_wick", "Liquidity taken by wick", "▼ biggest level", 17, "ict",
        "Biggest liquidity level the first spike (wick) took.",
        "▼ The biggest one taken: Prev week > Prev day > London > Asia > Pre-news > None.",
        "Pre-news high/low", LIQ),
    Col("liq_main", "Liquidity taken by main move (by 8:45)", "▼ biggest level", 19, "ict",
        "Biggest liquidity level the main move took by 8:45.",
        "▼ The biggest one taken: Prev week > Prev day > London > Asia > Pre-news > None.",
        "London high/low", LIQ),
    Col("pda", "HTF PD array at extreme (optional)", "▼ pick", 16, "ict",
        "Did the main move's extreme (by 8:45) tap a 15m / 1H / 4H PD array?",
        "▼ No / FVG / Order block / Breaker / Other.", "FVG", PDA),
    Col("fvg", "1m FVG from main move", "▼ pick", 24, "ict",
        "The first 1m fair value gap the main move left - what happened the first time price came back to it?",
        "▼ No FVG · Never came back (by 16:00) · Came back & held (no 1m close beyond its far edge) · "
        "Came back & failed (IFVG = a 1m candle closed beyond its far edge).", "Came back & held", FVG),
    Col("smt", "SMT with ES (optional)", "▼ Yes / No", 12, "ict",
        "At the main-move extreme, did NQ make a new high/low while ES didn't (or the other way round)?",
        "▼ Yes / No / Not checked.", "Not checked", SMT),
    # 6 · trade & review
    Col("setup", "Setup", "▼ pick", 18, "trade",
        "Which of your setups triggered, using fixed rules (see SETUPS TO TEST).",
        "▼ None / A Continuation / B Reversal / C Open sweep / D 9:30-11:00 raid.", "A Continuation", SETUP),
    Col("r", "Result (R)", "e.g. 2 or -1", 10, "trade",
        "Result of that setup in R (multiples of your risk).",
        "Number: 2 = +2R · -1 = full stop · 0 = breakeven. Leave empty if no setup.", 2,
        number_format="0.0#", validation="r"),
    Col("link", "Chart link", "TradingView link", 26, "trade",
        "TradingView snapshot of the day, so you can see it again in one click.",
        "Camera icon → Copy link → paste here.", "tradingview.com/x/…"),
]
for _i, _c in enumerate(COLUMNS, start=1):
    _c.letter = get_column_letter(_i)
C = {c.key: c.letter for c in COLUMNS}


def log_formula(key: str, r: int) -> str:
    """Formula for an auto column in row r of the log."""
    if key == "day":
        return f'=IF(ISNUMBER({C["date"]}{r}),TEXT({C["date"]}{r},"ddd"),"")'
    main, cpi = f'{C["main"]}{r}', f'{C["cpi"]}{r}'
    if key == "move_vs_data":
        return (f'=IF(OR({cpi}="",AND({main}<>"Up",{main}<>"Down")),"",'
                f'IF(OR({cpi}="In line",{cpi}="Mixed"),"{MOVE_VS_DATA[2]}",'
                f'IF(OR(AND({cpi}="Hotter",{main}="Down"),AND({cpi}="Cooler",{main}="Up")),'
                f'"{MOVE_VS_DATA[0]}","{MOVE_VS_DATA[1]}")))')
    hi, lo = f'{C["high_when"]}{r}', f'{C["low_when"]}{r}'
    if key == "outcome":
        pb = f'{C["pullback"]}{r}'
        mw = f'IF({main}="Up",{hi},{lo})'  # "taken - when" of the main-move side
        return (f'=IF(OR(AND({main}<>"Up",{main}<>"Down"),NOT(ISNUMBER({pb})),{mw}=""),"",'
                f'IF({pb}>=100,"{OUTCOMES[1]}",'
                f'IF(OR({mw}="{WIN_HL[0]}",{mw}="{WIN_HL[1]}"),"{OUTCOMES[0]}",'
                f'IF({pb}>=50,"{OUTCOMES[2]}","{OUTCOMES[3]}"))))')
    if key == "first_side":
        order = "|" + "|".join(WIN_HL) + "|"  # position in this string = time order
        p_hi, p_lo = f'FIND("|"&{hi}&"|","{order}")', f'FIND("|"&{lo}&"|","{order}")'
        high_first = f'IF({main}="Up","{HIGH_MAIN}",IF({main}="Down","{HIGH_WICK}","High first"))'
        low_first = f'IF({main}="Down","{LOW_MAIN}",IF({main}="Up","{LOW_WICK}","Low first"))'
        return (f'=IF(OR({hi}="",{lo}=""),"",IFERROR('
                f'IF(AND({hi}="{NEVER}",{lo}="{NEVER}"),"{NEITHER}",'
                f'IF({p_hi}={p_lo},"{SAME_WINDOW}",IF({p_hi}<{p_lo},{high_first},{low_first}))),'
                f'"Check entries"))')
    raise KeyError(key)


def style_font(cell, size=10, bold=False, italic=False, color="111827"):
    cell.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


# ---------------------------------------------------------------- CPI Log
def build_log(ws, demo_rows=None):
    ws.sheet_properties.tabColor = "111827"
    groups = []  # [section, first letter, last letter, text]
    for c in COLUMNS:
        if groups and groups[-1][0] == c.section:
            groups[-1][2] = c.letter
        else:
            groups.append([c.section, c.letter, c.letter, SECTIONS[c.section][0]])
    section_starts = {g[1] for g in groups}
    # Google Sheets can't freeze columns through a merged cell: split the band at the frozen edge.
    frozen_last, first_free = C["day"], C["wick"]
    legacy = groups[0]
    groups[0:1] = [["legacy", legacy[1], frozen_last, legacy[3]], ["legacy", first_free, legacy[2], LEGACY_BAND_REST]]
    for section, start, end, text in groups:
        band = SECTIONS[section][1]
        for col_idx in range(ws[f"{start}1"].column, ws[f"{end}1"].column + 1):
            ws.cell(row=1, column=col_idx).fill = fill(band)
        cell = ws[f"{start}1"]
        cell.value = text
        style_font(cell, 10, bold=text != LEGACY_BAND_REST, italic=text == LEGACY_BAND_REST, color="FFFFFF")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if start != end:
            ws.merge_cells(f"{start}1:{end}1")

    for c in COLUMNS:
        _, _, head, head_font = SECTIONS[c.section]
        edge = SECTION_EDGE if c.letter in section_starts else None
        ws.column_dimensions[c.letter].width = c.width
        h = ws[f"{c.letter}2"]
        h.value = c.header
        h.fill = fill(head)
        style_font(h, 10, bold=True, italic=c.auto, color=head_font)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = Border(bottom=THIN, left=edge or THIN, right=THIN)
        note = Comment(f"{c.header}\n\nWhat: {c.meaning}\n\nHow: {c.how}", "CPI template")
        note.width, note.height = 340, 200
        h.comment = note

        hint = ws[f"{c.letter}3"]
        hint.value = c.hint
        hint.fill = fill(AUTO_FILL if c.auto else "FFFFFF")
        style_font(hint, 8, italic=True, color="6B7280")
        hint.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hint.border = Border(bottom=Side(style="thin", color="9CA3AF"), left=edge)

        for r in range(FIRST_ROW, LAST_ROW + 1):
            cell = ws[f"{c.letter}{r}"]
            if c.auto:
                cell.value = log_formula(c.key, r)
                cell.fill = fill(AUTO_FILL)
                style_font(cell, 10, italic=True, color="374151")
            else:
                style_font(cell, 10)
            if c.number_format:
                cell.number_format = c.number_format
            if edge:
                cell.border = Border(left=edge)

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 54
    ws.row_dimensions[3].height = 34
    ws.freeze_panes = f"{first_free}{FIRST_ROW}"  # headers + date/day stay visible

    def col_range(key):
        return f"{C[key]}{FIRST_ROW}:{C[key]}{LAST_ROW}"

    for c in COLUMNS:
        if c.options:
            formula = '"' + ",".join(c.options) + '"'
            assert len(formula) <= 255 and all("," not in o for o in c.options), c.key
            dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                                errorStyle="stop", errorTitle=c.header[:32], error="Pick a value from the dropdown list.")
            dv.add(col_range(c.key))
            ws.add_data_validation(dv)
    pb = f"{C['pullback']}{FIRST_ROW}"
    dv = DataValidation(type="custom", formula1=f"AND(ISNUMBER({pb}),OR({pb}=0,AND({pb}>=1,{pb}<=1000)))",
                        allow_blank=True, showErrorMessage=True, errorStyle="stop", errorTitle="Deepest pullback %",
                        error="Type just the number, e.g. 62 for 62% (no % sign). 0 = no pullback, 100 = back to the 8:30 open.")
    dv.add(col_range("pullback"))
    ws.add_data_validation(dv)
    dv = DataValidation(type="decimal", operator="between", formula1="-50", formula2="50", allow_blank=True,
                        showErrorMessage=True, errorStyle="stop", errorTitle="Result (R)",
                        error="Type the result in R, e.g. 2 or -1 (a number between -50 and 50).")
    dv.add(col_range("r"))
    ws.add_data_validation(dv)

    cf = ws.conditional_formatting

    def equals(key, value, colour):
        cf.add(col_range(key), CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill(colour)))

    for key in ("wick", "main"):
        equals(key, "Up", GREEN)
        equals(key, "Down", RED)
    equals("straddle", "TP", GREEN)
    equals("straddle", "SL", RED)
    equals("straddle", "BE", AMBER)
    equals("move_vs_data", MOVE_VS_DATA[0], "D1FAE5")
    equals("move_vs_data", MOVE_VS_DATA[1], "FFEDD5")
    for label, colour in zip(OUTCOMES, (GREEN, RED, AMBER, GREY)):
        equals("outcome", label, colour)
    fs = f"{C['first_side']}{FIRST_ROW}"
    cf.add(col_range("first_side"), FormulaRule(formula=[f'ISNUMBER(SEARCH("main side",{fs}))'], fill=fill(GREEN)))
    cf.add(col_range("first_side"), FormulaRule(formula=[f'ISNUMBER(SEARCH("wick side",{fs}))'], fill=fill(RED)))
    cf.add(col_range("r"), CellIsRule(operator="greaterThan", formula=["0"], fill=fill(GREEN)))
    cf.add(col_range("r"), CellIsRule(operator="lessThan", formula=["0"], fill=fill(RED)))

    for offset, values in enumerate(demo_rows or []):
        for key, value in values.items():
            ws[f"{C[key]}{FIRST_ROW + offset}"].value = value

    ws.page_setup.orientation = "landscape"


# ---------------------------------------------------------------- Stats
def rng(key: str) -> str:
    return f"'{LOG}'!${C[key]}${FIRST_ROW}:${C[key]}${STATS_LAST_ROW}"


def q(text: str) -> str:
    return '"' + text.replace('"', '""') + '"'


DATE_FILTER = f'{rng("date")},">="&$C$3,{rng("date")},"<="&$E$3'


def cnt(*pairs) -> str:
    """COUNTIFS limited to the Stats date window. pairs = (column key, criterion expression)."""
    return "COUNTIFS(" + ",".join([DATE_FILTER] + [f"{rng(k)},{crit}" for k, crit in pairs]) + ")"


def avg(value_key: str, *pairs) -> str:
    crits = [DATE_FILTER] + [f"{rng(k)},{crit}" for k, crit in pairs]
    return f"AVERAGEIFS({rng(value_key)}," + ",".join(crits) + ")"


def pct(numerator: str, denominator: str) -> str:
    return f'=IFERROR(({numerator})/({denominator}),"–")'


class StatsWriter:
    """Writes stacked tables. Every formula cell is registered in `expect` with a spec the verifier recomputes."""

    def __init__(self, ws, expect):
        self.ws = ws
        self.row = 6
        self.expect = expect

    def put(self, cell, value, fmt=None, check=None, bold=False, fill_colour=None, align="center"):
        c = self.ws[cell]
        c.value = value
        style_font(c, 10, bold=bold)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if fmt:
            c.number_format = fmt
        if fill_colour:
            c.fill = fill(fill_colour)
        c.border = Border(bottom=THIN)
        if check is not None:
            self.expect[cell] = check

    def title(self, text, colour, note=None):
        ws = self.ws
        ws.merge_cells(f"A{self.row}:G{self.row}")
        cell = ws[f"A{self.row}"]
        cell.value = text
        cell.fill = fill(colour)
        style_font(cell, 11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[self.row].height = 20
        self.row += 1
        if note:
            ws.merge_cells(f"A{self.row}:G{self.row}")
            n = ws[f"A{self.row}"]
            n.value = note
            style_font(n, 9, italic=True, color="4B5563")
            n.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            ws.row_dimensions[self.row].height = 26
            self.row += 1

    def header(self, labels, colour):
        for i, label in enumerate(labels):
            self.put(f"{get_column_letter(i + 1)}{self.row}", label, bold=True, fill_colour=colour,
                     align="left" if i == 0 else "center")
        self.ws.row_dimensions[self.row].height = 30
        self.row += 1

    def label(self, text, bold=False):
        self.put(f"A{self.row}", text, align="left", bold=bold)


def build_stats(ws, expect=None):
    expect = {} if expect is None else expect
    ws.sheet_properties.tabColor = "047857"
    ws.column_dimensions["A"].width = 44
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 15
    ws.merge_cells("A1:G1")
    ws["A1"].value = "CPI $NQ – STATS (updates automatically from the CPI Log tab)"
    style_font(ws["A1"], 14, bold=True)
    ws.merge_cells("A2:G2")
    ws["A2"].value = ("Only rows you have filled are counted. '–' = no data yet. "
                      "Small samples lie: look at n before trusting a %.")
    style_font(ws["A2"], 9, italic=True, color="4B5563")
    ws["A3"].value = "Count releases from"
    ws["C3"].value = dt.datetime(2020, 1, 1)
    ws["D3"].value = "to"
    ws["E3"].value = "=TODAY()"
    ws["F3"].value = "← change these dates to test a period (e.g. 2023 onwards)"
    style_font(ws["A3"], 10, bold=True)
    style_font(ws["D3"], 10, bold=True)
    style_font(ws["F3"], 10, italic=True, color="4B5563")
    for cell in ("C3", "E3"):
        ws[cell].number_format = "dd mmm yyyy"
        ws[cell].fill = fill("FEF9C3")
        style_font(ws[cell], 10, bold=True)
        ws[cell].alignment = Alignment(horizontal="center")
    ws["D3"].alignment = Alignment(horizontal="center")
    ws["A4"].value = "Releases in this period"
    style_font(ws["A4"], 10, bold=True)
    ws["C4"].value = f"=COUNTIFS({DATE_FILTER})"
    expect["C4"] = ("releases",)
    ws["D4"].value = "with an Outcome:"
    ws["E4"].value = "=" + "+".join(cnt(("outcome", q(o))) for o in OUTCOMES)
    expect["E4"] = ("n", "outcome", OUTCOMES, None, None)
    for cell, align in (("C4", "center"), ("D4", "right"), ("E4", "center")):
        style_font(ws[cell], 10, bold=cell != "D4")
        ws[cell].alignment = Alignment(horizontal=align)
    ws.freeze_panes = "A5"

    s = StatsWriter(ws, expect)
    P, N, D1, D2 = "0%", "0", "0.0", "0.00"

    def count_table(key, labels, first_col_title, share_of=None, cumulative=False, extra_share=None):
        """# and % for each label of one column (+ optional cumulative % / sub-share)."""
        heads = [first_col_title, "#", "%"]
        if cumulative:
            heads.append("Cumulative %")
        if extra_share:
            heads.append(extra_share[0])
        s.header(heads, s.head_colour)
        first, n_row = s.row, s.row + len(labels)
        for label in labels:
            s.label(label)
            s.put(f"B{s.row}", "=" + cnt((key, f"$A{s.row}")), N, ("count", key, label))
            s.put(f"C{s.row}", f'=IFERROR(B{s.row}/B${n_row},"–")', P, ("share", key, [label], labels))
            if cumulative and label != NEVER:
                upto = labels[:labels.index(label) + 1]
                s.put(f"D{s.row}", f'=IFERROR(SUM(B${first}:B{s.row})/B${n_row},"–")', P,
                      ("share", key, upto, labels))
            if extra_share and label in extra_share[1]:
                sub = extra_share[1]
                lo, hi = first + labels.index(sub[0]), first + labels.index(sub[-1])
                s.put(f"D{s.row}", f'=IFERROR(B{s.row}/SUM(B${lo}:B${hi}),"–")', P, ("share", key, [label], sub))
            s.row += 1
        s.label(share_of or "n (releases)", bold=True)
        s.put(f"B{s.row}", f"=SUM(B{first}:B{s.row - 1})", N, ("total", key, labels), bold=True)
        s.row += 3

    # 1 — outcome after the first minute, split by CPI result
    s.title("1 · What did the main move do after the 1st minute? (Outcome 8:31–8:45)", "047857",
            "Read across: e.g. 'when CPI was Hotter, the main move went back to the 8:30 open X% of the time'.")
    s.header(["Outcome", "All", "Hotter", "Cooler", "In line", "Mixed"], "D1FAE5")
    n_row = s.row + len(OUTCOMES)
    groups = [None] + CPI
    for o in OUTCOMES:
        s.label(o)
        for i, cpi in enumerate(groups):
            col = "BCDEF"[i]
            pairs = [("outcome", f"$A{s.row}")] + ([("cpi", q(cpi))] if cpi else [])
            s.put(f"{col}{s.row}", pct(cnt(*pairs), f"{col}${n_row}"), P, ("cond_share", "cpi", cpi, o))
        s.row += 1
    s.label("n (releases)", bold=True)
    for i, cpi in enumerate(groups):
        extra = [("cpi", q(cpi))] if cpi else []
        s.put(f"{'BCDEF'[i]}{s.row}", "=" + "+".join(cnt(("outcome", q(o)), *extra) for o in OUTCOMES), N,
              ("n", "outcome", OUTCOMES, "cpi", cpi), bold=True)
    assert s.row == n_row
    s.row += 3

    # 2 / 3 — back to the open
    s.head_colour = "D1FAE5"
    s.title("2 · When did price come back to the 8:30 open?", "047857",
            "Cumulative % = it had come back by the end of that window.")
    count_table("back_open", WIN_OPEN, "Window (NY time)", cumulative=True)
    s.title("3 · At the 8:30 open, price…  (does it just grab the liquidity and keep going?)", "047857")
    count_table("at_open", AT_OPEN, "Reaction", extra_share=("% of the times it came back", AT_OPEN[:3]))

    # 4 — 1-min high / low taken
    s.title("4 · How long until the 1-min candle's HIGH / LOW (BSL / SSL) got taken?", "047857",
            "Main side = the end of the 1-min candle in the Main Move direction (the HIGH when Main Move is Up). "
            "Wick side = the other end. Cumulative % = taken by the end of that window. "
            "'Not by 16:00' = that end held as the high/low for the rest of the day.")
    s.header(["Window (NY time)", "Main side #", "Main side cum. %", "Wick side #", "Wick side cum. %"], "D1FAE5")
    first, n_row = s.row, s.row + len(WIN_HL)
    for w in WIN_HL:
        s.label(w)
        a = f"$A{s.row}"
        main_side = f'{cnt(("main", q("Up")), ("high_when", a))}+{cnt(("main", q("Down")), ("low_when", a))}'
        wick_side = f'{cnt(("main", q("Up")), ("low_when", a))}+{cnt(("main", q("Down")), ("high_when", a))}'
        s.put(f"B{s.row}", "=" + main_side, N, ("side", "main", [w]))
        s.put(f"D{s.row}", "=" + wick_side, N, ("side", "wick", [w]))
        if w != NEVER:
            upto = WIN_HL[:WIN_HL.index(w) + 1]
            s.put(f"C{s.row}", f'=IFERROR(SUM(B${first}:B{s.row})/B${n_row},"–")', P, ("side_share", "main", upto))
            s.put(f"E{s.row}", f'=IFERROR(SUM(D${first}:D{s.row})/D${n_row},"–")', P, ("side_share", "wick", upto))
        s.row += 1
    s.label("n (releases)", bold=True)
    s.put(f"B{s.row}", f"=SUM(B{first}:B{s.row - 1})", N, ("side", "main", WIN_HL), bold=True)
    s.put(f"D{s.row}", f"=SUM(D{first}:D{s.row - 1})", N, ("side", "wick", WIN_HL), bold=True)
    s.row += 2
    s.header(["Which side was taken first?", "#", "%"], "D1FAE5")
    first_side_rows = [
        ("Main side first (continuation)", [HIGH_MAIN, LOW_MAIN]),
        ("Wick side first (reversal)", [HIGH_WICK, LOW_WICK]),
        ("Both in the same window", [SAME_WINDOW]),
        ("Neither taken by 16:00", [NEITHER]),
    ]
    all_first = [l for _, ls in first_side_rows for l in ls]
    first, n_row = s.row, s.row + len(first_side_rows)
    for text, labels in first_side_rows:
        s.label(text)
        s.put(f"B{s.row}", "=" + "+".join(cnt(("first_side", q(l))) for l in labels), N,
              ("total", "first_side", labels))
        s.put(f"C{s.row}", f'=IFERROR(B{s.row}/B${n_row},"–")', P, ("share", "first_side", labels, all_first))
        s.row += 1
    s.label("n (releases)", bold=True)
    s.put(f"B{s.row}", f"=SUM(B{first}:B{s.row - 1})", N, ("total", "first_side", all_first), bold=True)
    s.row += 3

    # 5 — candle close location relative to the main move
    s.title("5 · Where did the candles close? (relative to the main move)", "6D28D9",
            "Strong end = the third on the main-move side (Top when Main Move is Up, Bottom when it is Down). "
            "Wick end = the opposite third.")
    s.header(["Candle", "Strong end", "Middle", "Wick end", "n"], "EDE9FE")
    for text, key in (("1m candle (8:30)", "close1"), ("5m candle (8:30-8:34)", "close5"),
                      ("15m candle (8:30-8:44)", "close15")):
        s.label(text)
        strong = f'{cnt(("main", q("Up")), (key, q("Top")))}+{cnt(("main", q("Down")), (key, q("Bottom")))}'
        middle = f'{cnt(("main", q("Up")), (key, q("Middle")))}+{cnt(("main", q("Down")), (key, q("Middle")))}'
        weak = f'{cnt(("main", q("Up")), (key, q("Bottom")))}+{cnt(("main", q("Down")), (key, q("Top")))}'
        s.put(f"E{s.row}", f"={strong}+{middle}+{weak}", N, ("close", key, None), bold=True)
        s.put(f"B{s.row}", pct(strong, f"E{s.row}"), P, ("close", key, "strong"))
        s.put(f"C{s.row}", pct(middle, f"E{s.row}"), P, ("close", key, "middle"))
        s.put(f"D{s.row}", pct(weak, f"E{s.row}"), P, ("close", key, "weak"))
        s.row += 1
    s.row += 2

    # 6 — pullback depth
    s.title("6 · How deep was the pullback before the main move continued?", "047857",
            "Use this for entries and stops: e.g. if most Continued cases reach 62% (OTE), "
            "a limit order at 62% gets filled most of the time.")
    s.header(["Deepest pullback %", "All", "When it Continued"], "D1FAE5")
    n_row = s.row + 6
    cont = ("outcome", q(OUTCOMES[0]))
    s.label("Average pullback")
    s.put(f"B{s.row}", f'=IFERROR({avg("pullback", ("pullback", q(">=0")))},"–")', '0"%"', ("pb_avg", None))
    s.put(f"C{s.row}", f'=IFERROR({avg("pullback", cont, ("pullback", q(">=0")))},"–")', '0"%"',
          ("pb_avg", OUTCOMES[0]))
    s.row += 1
    for text, level in (("Reached 50% (half-way / equilibrium)", 50), ("Reached 62% (OTE starts)", 62),
                        ("Reached 70.5% (OTE sweet spot)", 70.5), ("Reached 79% (OTE ends)", 79),
                        ("Reached 100% (back to the 8:30 open)", 100)):
        s.label(text)
        s.put(f"B{s.row}", pct(cnt(("pullback", q(f">={level}"))), f"B${n_row}"), P, ("pb_reach", level, None))
        if level < 100:  # Continued means the pullback stayed under 100 by definition
            s.put(f"C{s.row}", pct(cnt(("pullback", q(f">={level}")), cont), f"C${n_row}"), P,
                  ("pb_reach", level, OUTCOMES[0]))
        s.row += 1
    s.label("n (releases)", bold=True)
    s.put(f"B{s.row}", "=" + cnt(("pullback", q(">=0"))), N, ("pb_n", None), bold=True)
    s.put(f"C{s.row}", "=" + cnt(("pullback", q(">=0")), cont), N, ("pb_n", OUTCOMES[0]), bold=True)
    assert s.row == n_row
    s.row += 3

    # 7 — CPI number vs direction
    s.head_colour = "DBEAFE"
    s.title("7 · Did NQ move the 'logical' way for the CPI number?", "1D4ED8",
            "With data = Hotter → Main Move Down, or Cooler → Main Move Up. In line / Mixed are left out.")
    count_table("move_vs_data", MOVE_VS_DATA[:2], "Main move vs CPI", share_of="n (Hotter or Cooler)")
    s.row -= 1
    s.header(["CPI vs forecast", "Releases", "Main move Up", "Main move Down"], "DBEAFE")
    for cpi in CPI:
        s.label(cpi)
        up, down = cnt(("cpi", f"$A{s.row}"), ("main", q("Up"))), cnt(("cpi", f"$A{s.row}"), ("main", q("Down")))
        s.put(f"B{s.row}", f"={up}+{down}", N, ("cpi_dir", cpi, None))
        s.put(f"C{s.row}", pct(up, f"B{s.row}"), P, ("cpi_dir", cpi, "Up"))
        s.put(f"D{s.row}", pct(down, f"B{s.row}"), P, ("cpi_dir", cpi, "Down"))
        s.row += 1
    s.row += 2

    # 8 — averages of the original columns
    s.title("8 · Average size in points (your original columns)", "374151",
            "Only real numbers are averaged - cells with text like 'Same' are skipped.")
    s.header(["Average (points)", "All", "Main move Up", "Main move Down"], "E5E7EB")
    for text, key in (("Wick pts (when there was a wick)", "wick_pts"), ("15sec move", "m15"),
                      ("30sec move", "m30"), ("1min move", "m1"), ("Scam Wick (15 sec before)", "scam")):
        s.label(text)
        num = (key, q(">-1000000"))
        s.put(f"B{s.row}", f'=IFERROR({avg(key, num)},"–")', D1, ("avg", key, None))
        s.put(f"C{s.row}", f'=IFERROR({avg(key, num, ("main", q("Up")))},"–")', D1, ("avg", key, "Up"))
        s.put(f"D{s.row}", f'=IFERROR({avg(key, num, ("main", q("Down")))},"–")', D1, ("avg", key, "Down"))
        s.row += 1
    s.row += 2

    # 9 — straddle
    s.head_colour = "E5E7EB"
    s.title("9 · Straddle results", "374151",
            "Chart backtests don't show news slippage/spread - real fills on stop orders are usually worse.")
    count_table("straddle", STRADDLE[:3], "Result", share_of="n (N/A not counted)")

    # 10 / 11 / 12 — ICT conditions → outcome
    def outcome_rows(first_col_title, rows):
        s.header([first_col_title, "n"] + OUTCOMES, "FEE2E2")
        for text, key, value in rows:
            s.label(text)
            crit = f"$A{s.row}" if text == value else q(value)
            s.put(f"B{s.row}", "=" + "+".join(cnt((key, crit), ("outcome", q(o))) for o in OUTCOMES), N,
                  ("n", "outcome", OUTCOMES, key, value), bold=True)
            for i, o in enumerate(OUTCOMES):
                col = "CDEF"[i]
                s.put(f"{col}{s.row}", pct(cnt((key, crit), ("outcome", q(o))), f"$B{s.row}"), P,
                      ("cond_share", key, value, o))
            s.row += 1
        s.row += 2

    s.title("10 · ICT: liquidity taken by the MAIN MOVE (by 8:45) → what happened after the 1st minute?", "B91C1C")
    outcome_rows("Biggest level taken", [(level, "liq_main", level) for level in LIQ])
    s.title("11 · ICT: liquidity taken by the WICK → what happened after the 1st minute?", "B91C1C")
    outcome_rows("Biggest level taken", [(level, "liq_wick", level) for level in LIQ])
    s.title("12 · ICT: conditions → what happened after the 1st minute?", "B91C1C",
            "A condition is useful when its % is clearly different from the 'All' column of table 1.")
    outcome_rows("Condition",
                 [(f"Daily bias: {b}", "bias", b) for b in BIAS]
                 + [("8:30 open above Midnight open", "midnight", MIDNIGHT[0]),
                    ("8:30 open below Midnight open", "midnight", MIDNIGHT[1])]
                 + [(f"HTF PD array at extreme: {p}", "pda", p) for p in PDA]
                 + [(f"1m FVG: {f}", "fvg", f) for f in FVG]
                 + [(f"SMT with ES: {m}", "smt", m) for m in SMT[:2]])
    s.row -= 1
    s.header(["Did the main move follow your daily bias?", "#", "%"], "FEE2E2")
    first, n_row = s.row, s.row + 2
    with_bias = f'{cnt(("bias", q("Bullish")), ("main", q("Up")))}+{cnt(("bias", q("Bearish")), ("main", q("Down")))}'
    against = f'{cnt(("bias", q("Bullish")), ("main", q("Down")))}+{cnt(("bias", q("Bearish")), ("main", q("Up")))}'
    for text, formula, kind in (("Yes (Bullish → Up, Bearish → Down)", with_bias, "with"),
                                ("No (went against the bias)", against, "against")):
        s.label(text)
        s.put(f"B{s.row}", "=" + formula, N, ("bias", kind))
        s.put(f"C{s.row}", f'=IFERROR(B{s.row}/B${n_row},"–")', P, ("bias", kind + "_share"))
        s.row += 1
    s.label("n (days with a bias)", bold=True)
    s.put(f"B{s.row}", f"=SUM(B{first}:B{s.row - 1})", N, ("bias", "n"), bold=True)
    s.row += 3

    # 13 — setups
    s.title("13 · Your setups (only rows with a Result in R)", "B45309",
            "Avg R = expectancy. A setup is worth trading when Avg R stays positive over a decent n (30+).")
    s.header(["Setup", "Trades", "Wins (R > 0)", "Win rate", "Avg R", "Total R"], "FEF3C7")
    first = s.row
    for setup in SETUP[1:]:
        a = f"$A{s.row}"
        s.label(setup)
        s.put(f"B{s.row}", "=" + cnt(("setup", a), ("r", q(">-1000"))), N, ("setup", setup, "trades"))
        s.put(f"C{s.row}", "=" + cnt(("setup", a), ("r", q(">0"))), N, ("setup", setup, "wins"))
        s.put(f"D{s.row}", f'=IFERROR(C{s.row}/B{s.row},"–")', P, ("setup", setup, "winrate"))
        s.put(f"E{s.row}", f'=IFERROR({avg("r", ("setup", a), ("r", q(">-1000")))},"–")', D2,
              ("setup", setup, "avg"))
        s.put(f"F{s.row}", f'=SUMIFS({rng("r")},{DATE_FILTER},{rng("setup")},{a})', D2, ("setup", setup, "total"))
        s.row += 1
    last = s.row - 1
    s.label("All setups", bold=True)
    s.put(f"B{s.row}", f"=SUM(B{first}:B{last})", N, ("setup", None, "trades"), bold=True)
    s.put(f"C{s.row}", f"=SUM(C{first}:C{last})", N, ("setup", None, "wins"), bold=True)
    s.put(f"D{s.row}", f'=IFERROR(C{s.row}/B{s.row},"–")', P, ("setup", None, "winrate"), bold=True)
    s.put(f"E{s.row}", f'=IFERROR(F{s.row}/B{s.row},"–")', D2, ("setup", None, "avg"), bold=True)
    s.put(f"F{s.row}", f"=SUM(F{first}:F{last})", D2, ("setup", None, "total"), bold=True)

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return expect


# ---------------------------------------------------------------- Guide
QUICK_START = [
    ("One row = one CPI", "Work left → right while you replay the day with TradingView Bar Replay. "
                          "The columns follow the order things happen."),
    ("New York time", "Set the TradingView chart timezone to New York (UTC-4/-5). CPI is always 8:30 NY time "
                      "and every time window in this sheet uses NY time."),
    ("White vs grey", "White cells: you fill them. Grey cells: automatic - don't type in them."),
    ("Dropdowns & numbers", "▼ = pick from the dropdown. Number columns take numbers only "
                            "(instead of 'Same', type the number again)."),
    ("No hindsight", "Fill the ICT context (bias, levels to watch) at 8:29, BEFORE you play the release."),
    ("Chart link", "For every row save a TradingView snapshot (camera icon → Copy link) in 'Chart link', "
                   "so later you can see the day in one click."),
    ("Stats", "The Stats tab updates by itself. To test only a period (e.g. 2023+), change the two dates "
              "at the top of Stats."),
]

WORDS = [
    ("8:30 open", "The open price of the 8:30:00 candle = the price at the moment CPI comes out."),
    ("1-min candle", "The 8:30 one-minute candle (8:30:00-8:30:59). Its HIGH is buy-side liquidity (BSL), "
                     "its LOW is sell-side liquidity (SSL). Your 'initial wick' high/low is inside this candle."),
    ("Main side / Wick side", "Main side = the end of the 1-min candle in your Main Move direction "
                              "(the HIGH when Main Move is Up). Wick side = the other end."),
    ("Taken", "Price trades beyond the level, even by 1 tick. A sweep counts."),
    ("Time windows (NY)", "0-15 sec = 8:30:00-8:30:14 · 15-30 sec = 8:30:15-8:30:29 · 30-60 sec = 8:30:30-8:30:59 · "
                          "1-5 min = 8:31:00-8:34:59 · 5-15 min = 8:35:00-8:44:59 · 15-60 min = 8:45:00-9:29:59 · "
                          "9:30-11:00 · 11:00-16:00 · Not by 16:00 = didn't happen that day."),
    ("Deepest pullback %", "Put a Fib with 0 on the 1-min main-side extreme and 100 on the 8:30 open "
                           "(tick 'Reverse' in the Fib settings if it's upside down). Record the deepest level "
                           "price reached BEFORE it broke the 1-min main-side extreme; stop checking at 8:45. "
                           "0 = no pullback · 50 = half-way · 62-79 = OTE · 100 = back at the 8:30 open · "
                           "120 = 20% past it. Type 62, not 62%."),
    ("Outcome (auto)", "Back to open = pullback reached 100 · Continued = main side taken by 8:45 "
                       "(1-5 or 5-15 min) with pullback under 100 · Partial retrace = not taken by 8:45, "
                       "pullback 50-99 · Stalled / chop = not taken by 8:45, pullback under 50."),
    ("Candle close", "Split the candle's full range (wicks included) into 3 equal parts: which part did it close "
                     "in? 1m = 8:30 candle · 5m = 8:30-8:34 · 15m = 8:30-8:44."),
    ("Back to 8:30 open", "The first time price trades back at the 8:30 open after the main move started."),
    ("At the open, price…", "Swept & resumed = after touching the open it turned and made a new high/low in the "
                            "main direction before taking the 1-min wick side · Broke through = it took the 1-min "
                            "wick side first · Chopped around = neither by 16:00 · Never came back."),
    ("CPI vs forecast", "ForexFactory history: CPI m/m and Core CPI m/m, actual vs forecast. Both higher (or one "
                        "higher, one equal) = Hotter · both lower (or one lower, one equal) = Cooler · both equal = "
                        "In line · one higher, one lower = Mixed."),
    ("Liquidity levels", "Pre-news = high/low of 7:00-8:29 · London = 2:00-5:00 · Asia = 20:00-00:00 (evening "
                         "before) · Prev day = yesterday's daily candle · Prev week = last week's candle. "
                         "Pick the BIGGEST one taken: Prev week > Prev day > London > Asia > Pre-news. "
                         "None = nothing taken."),
    ("Midnight open", "Open of the 00:00 New York candle. 8:30 open above it = premium, below it = discount."),
    ("HTF PD array", "Fair value gap, order block or breaker on 15m / 1H / 4H."),
    ("1m FVG", "The first fair value gap the main move left on the 1m chart. Held = price came back into it and "
               "bounced with no 1m close beyond its far edge. Failed (IFVG) = a 1m candle closed beyond its far "
               "edge."),
    ("SMT", "At the main-move extreme NQ made a new high/low and ES did not (or the other way round)."),
    ("R", "Result in multiples of your risk: +2 = made 2x the risk, -1 = full stop loss, 0 = breakeven."),
]

SETUPS = [
    ("A Continuation", "Main move leaves a 1m FVG → enter on the first pullback into it (Stats table 6 shows how "
                       "deep pullbacks usually go) → stop beyond the pullback's swing → target the 1-min main "
                       "side, then the next liquidity."),
    ("B Reversal", "Main move takes a key level (London / prev day) or taps an HTF PD array, ideally with SMT → "
                   "1m MSS back through the last swing + IFVG → enter on the retrace → stop beyond the extreme → "
                   "target the 8:30 open, then the 1-min wick side."),
    ("C Open sweep", "Price comes back to the 8:30 open, sweeps it, then shifts back in the main direction "
                     "(1m MSS) → enter → stop beyond the sweep → target the 1-min main side."),
    ("D 9:30-11:00 raid", "After 9:30 price takes one side of the 1-min candle, then MSS + FVG → enter → "
                          "target the other side."),
    ("Rules first", "Write exact entry / stop / target rules for each setup BEFORE you backtest and never change "
                    "them mid-test. Log every trigger, winners and losers. No valid trigger = Setup 'None'."),
]

ROUTINE = [
    ("1 · 8:29 (pause)", "Mark Asia, London and pre-news (7:00-8:29) highs/lows, previous day/week high/low, "
                         "the Midnight open and any 15m/1H/4H FVG or order block near price.",
     "Daily bias · 8:30 open vs Midnight open"),
    ("2 · 8:30:00-8:30:59", "15s chart: first spike, main move, how far in 15s/30s/1m, straddle, scam wick. "
                            "Which level did the wick take? Did the main move leave a 1m FVG?",
     "Your columns · Liquidity taken by wick"),
    ("3 · 8:31-8:45", "1m chart: candle closes, deepest pullback with the Fib, which level the main move took, "
                      "PD array at the extreme, SMT with ES. Look up the CPI actual vs forecast.",
     "CPI vs forecast · Candle closes · Deepest pullback % · Liquidity by main move · PD array · SMT"),
    ("4 · until 16:00", "When the 1-min HIGH and LOW got taken, when price came back to the 8:30 open and what it "
                        "did there, and what happened when price came back to the 1m FVG.",
     "HIGH/LOW taken · Back to open · At the open · 1m FVG"),
    ("5 · Review", "Did a setup trigger by your rules? Record it and the result in R, save the chart link, "
                   "and check the grey auto columns make sense.",
     "Setup · Result (R) · Chart link"),
]


def build_guide(ws):
    ws.sheet_properties.tabColor = "1D4ED8"
    widths = {"A": 6, "B": 26, "C": 78, "D": 52, "E": 22}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    row = 2

    def section(text, colour):
        nonlocal row
        row += 1
        ws.merge_cells(f"A{row}:E{row}")
        cell = ws[f"A{row}"]
        cell.value = text
        cell.fill = fill(colour)
        style_font(cell, 11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 20
        row += 1

    def line(values, fills=None, bold_all=False):
        nonlocal row
        tallest = 1
        for i, value in enumerate(values):
            col = "ABCDE"[i]
            cell = ws[f"{col}{row}"]
            cell.value = value
            style_font(cell, 10, bold=bold_all or (i <= 1 and value is not None))
            cell.alignment = top
            cell.border = Border(bottom=THIN)
            if fills and fills[i]:
                cell.fill = fill(fills[i])
            if isinstance(value, str):
                # Imported rows don't always auto-grow, so size them from the text length.
                per_line = max(6, widths[col] * 1.05)
                tallest = max(tallest, math.ceil(len(value) / per_line))
        ws.row_dimensions[row].height = max(16, 13 * tallest + 5)
        row += 1

    ws.merge_cells("A1:E1")
    ws["A1"].value = "HOW TO USE THIS CPI BACKTEST SHEET"
    style_font(ws["A1"], 16, bold=True)
    ws.row_dimensions[1].height = 26
    ws.merge_cells("A2:E2")
    ws["A2"].value = "Hover any header in the CPI Log tab to see the same definitions there."
    style_font(ws["A2"], 10, italic=True, color="4B5563")

    section("QUICK START", "1D4ED8")
    for i, (title, text) in enumerate(QUICK_START, start=1):
        line([i, title, text])

    section("WORDS USED IN THIS SHEET", "1D4ED8")
    for term, text in WORDS:
        line([None, term, text])

    section("COLUMN GUIDE  (the Example column is made up - it shows one complete row)", "1D4ED8")
    line(["Col", "Column", "What it means", "How to fill", "Example"], fills=["E5E7EB"] * 5, bold_all=True)
    for c in COLUMNS:
        tint = SECTIONS[c.section][2] if c.section != "legacy" else "E5E7EB"
        example = "Wed (auto)" if c.key == "day" else c.example
        line([c.letter, c.header, c.meaning, c.how, example],
             fills=[tint, tint, None, AUTO_FILL if c.auto else None, None])

    section("SETUPS TO TEST", "B45309")
    for name, text in SETUPS:
        line([None, name, text])

    section("REPLAY ROUTINE (ICT checklist) - the columns follow this order", "B91C1C")
    line([None, "Step", "What to check", "Columns you fill"], fills=[None, "E5E7EB", "E5E7EB", "E5E7EB"],
         bold_all=True)
    for step, text, cols in ROUTINE:
        line([None, step, text, cols])
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


DEMO_ROWS = [
    dict(date="Example A", wick="Down", wick_pts=12, main="Up", m15=80, m30=95, m1=110, straddle="TP", scam=6,
         notes="EXAMPLE - made-up numbers", cpi="Cooler", close1="Top", close5="Top", close15="Middle",
         high_when="5-15 min", low_when=NEVER, pullback=62, back_open="15-60 min", at_open="Swept & resumed",
         bias="Bullish", midnight="Below (discount)", liq_wick="Pre-news high/low", liq_main="London high/low",
         pda="FVG", fvg="Came back & held", smt="Not checked", setup="A Continuation", r=2,
         link="tradingview.com/x/…"),
    dict(date="Example B", wick="Up", wick_pts=8, main="Down", m15=60, m30=75, m1=75, straddle="SL", scam=5,
         notes="EXAMPLE - made-up numbers", cpi="Hotter", close1="Bottom", close5="Middle", close15="Top",
         high_when="5-15 min", low_when=NEVER, pullback=130, back_open="1-5 min", at_open="Broke through",
         bias="Bullish", midnight="Below (discount)", liq_wick="None", liq_main="London high/low",
         pda="FVG", fvg="Came back & failed (IFVG)", smt="Yes", setup="B Reversal", r=3,
         link="tradingview.com/x/…"),
    dict(date="Example C", wick="No Wick", main="Up", m15=40, m30=45, m1=45, straddle="BE", scam=3,
         notes="EXAMPLE - made-up numbers", cpi="In line", close1="Middle", close5="Middle", close15="Middle",
         high_when="15-60 min", low_when="11:00-16:00", pullback=35, back_open="9:30-11:00",
         at_open="Broke through", bias="No bias", midnight="Above (premium)", liq_wick="None",
         liq_main="Pre-news high/low", pda="No", fvg="No FVG", smt="No", setup="None",
         link="tradingview.com/x/…"),
]


def build(path: Path, demo: bool = False, expect: dict | None = None) -> dict:
    wb = Workbook()
    log = wb.active
    log.title = LOG
    build_log(log, DEMO_ROWS if demo else None)
    expect = build_stats(wb.create_sheet(STATS), expect)
    build_guide(wb.create_sheet(GUIDE))
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return expect


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default="CPI_NQ_Backtest_Template.xlsx")
    parser.add_argument("--demo", action="store_true", help="add 3 made-up example rows (for screenshots)")
    args = parser.parse_args()
    build(Path(args.out), demo=args.demo)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
