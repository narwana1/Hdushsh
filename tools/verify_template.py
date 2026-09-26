#!/usr/bin/env python3
"""Recalculate the template with LibreOffice on test data and check every formula.

Fills the CPI Log with seeded synthetic rows (plus edge cases), lets LibreOffice
compute the workbook, then compares every auto column and every Stats cell with
an independent Python implementation of the same rules.

    python tools/verify_template.py                 # synthetic data only
    python tools/verify_template.py --legacy old.xlsx  # also paste A:K from an export of the old tab
"""
from __future__ import annotations

import argparse
import datetime as dt
import random
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_template as bt  # noqa: E402

LEGACY_KEYS = ["date", "day", "wick", "wick_pts", "main", "m15", "m30", "m1", "straddle", "scam", "notes"]
INPUT_OPTIONS = {c.key: c.options for c in bt.COLUMNS if c.options and c.section != "legacy"}
WINDOW = (dt.datetime(2021, 1, 1), dt.datetime(2025, 12, 31))
ERRORS = ("#VALUE!", "#NAME?", "#DIV/0!", "#REF!", "#N/A", "#NUM!", "Err:", "#NULL!")


def isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def blank(v):
    return v is None or v == ""


# ------------------------------------------------------------------ test data
def synthetic_rows(rng: random.Random, n: int):
    rows = []
    day = dt.datetime(2019, 12, 10)
    for i in range(n):
        day += dt.timedelta(days=rng.randint(25, 35))
        main = rng.choice(["Up", "Down", "Up", "Down", None])
        row = {
            "date": day,
            "wick": rng.choice(bt.WICK + [None]),
            "wick_pts": rng.choice([None, round(rng.uniform(0, 120), 2)]),
            "main": main,
            "m15": round(rng.uniform(0, 300), 2),
            "m30": rng.choice([round(rng.uniform(0, 300), 2), "Same", "26.5 opposite"]),
            "m1": rng.choice([round(rng.uniform(0, 400), 2), "Same", "N/A"]),
            "straddle": rng.choice(bt.STRADDLE + [None]),
            "scam": rng.choice([None, round(rng.uniform(0, 90), 2)]),
        }
        for key, options in INPUT_OPTIONS.items():
            row[key] = rng.choice(options + [None])
        row["pullback"] = rng.choice([None, 0, 12, 49.9, 50, 61.8, 62, 70.5, 79, 99.9, 100, 135, rng.uniform(0, 200)])
        row["r"] = rng.choice([None, -1, 0, 0.5, 1.5, 2, 3, -1.2]) if row["setup"] not in (None, "None") else None
        if i % 7 == 0:  # the user pastes old rows with the weekday typed in
            row["day"] = day.strftime("%a")
        rows.append(row)
    # edge cases
    rows.append({"date": "Canceled", "notes": "Canceled release"})
    rows.append({"date": dt.datetime(2022, 5, 11), "main": "Up", "pullback": 100, "high_when": "1-5 min",
                 "low_when": "1-5 min", "cpi": "Hotter", "at_open": "Swept & resumed"})
    rows.append({"date": dt.datetime(2022, 6, 10), "main": "Down", "pullback": 99.99, "high_when": bt.NEVER,
                 "low_when": "5-15 min", "cpi": "Cooler"})
    rows.append({"date": dt.datetime(2022, 7, 13), "main": "Down", "pullback": 0, "high_when": bt.NEVER,
                 "low_when": bt.NEVER, "cpi": "Mixed"})
    rows.append({"date": dt.datetime(2022, 8, 10), "main": None, "pullback": 40, "high_when": "1-5 min",
                 "low_when": "9:30-11:00"})
    rows.append({"date": dt.datetime(2020, 1, 14), "main": "Up", "pullback": 20, "high_when": "1-5 min",
                 "low_when": bt.NEVER, "setup": "A Continuation", "r": 2})  # outside the Stats window
    return rows


def legacy_rows(path: Path):
    ws = load_workbook(path).worksheets[0]
    out = []
    for values in ws.iter_rows(min_row=2, max_col=11, values_only=True):
        if all(blank(v) for v in values):
            continue
        out.append(dict(zip(LEGACY_KEYS, values)))
    return out


def add_random_new_columns(rows, rng):
    for row in rows:
        if not isinstance(row.get("date"), dt.datetime):
            continue
        for key, options in INPUT_OPTIONS.items():
            row[key] = rng.choice(options)
        row["pullback"] = rng.choice([0, 25, 50, 62, 79, 100, 120, rng.uniform(0, 160)])
        row["r"] = rng.choice([-1, 2, 3]) if row["setup"] != "None" else None


# ------------------------------------------------------------------ expected values
def derive(row):
    main, cpi = row.get("main"), row.get("cpi")
    d = {}
    date = row.get("date")
    d["day"] = row["day"] if not blank(row.get("day")) else (date.strftime("%a") if isinstance(date, dt.datetime) else "")
    if blank(cpi) or main not in ("Up", "Down"):
        d["move_vs_data"] = ""
    elif cpi in ("In line", "Mixed"):
        d["move_vs_data"] = "Neutral data"
    elif (cpi, main) in (("Hotter", "Down"), ("Cooler", "Up")):
        d["move_vs_data"] = "With data"
    else:
        d["move_vs_data"] = "Against data"

    hi, lo, pb = row.get("high_when"), row.get("low_when"), row.get("pullback")
    main_when = hi if main == "Up" else lo
    if main not in ("Up", "Down") or not isnum(pb) or blank(main_when):
        d["outcome"] = ""
    elif pb >= 100:
        d["outcome"] = "Back to open"
    elif main_when in ("1-5 min", "5-15 min"):
        d["outcome"] = "Continued"
    elif pb >= 50:
        d["outcome"] = "Partial retrace"
    else:
        d["outcome"] = "Stalled / chop"

    if blank(hi) or blank(lo):
        d["first_side"] = ""
    elif hi == bt.NEVER and lo == bt.NEVER:
        d["first_side"] = "Neither taken"
    else:
        ih, il = bt.WIN_HL.index(hi), bt.WIN_HL.index(lo)
        if ih == il:
            d["first_side"] = "Same window"
        elif ih < il:
            d["first_side"] = {"Up": bt.HIGH_MAIN, "Down": bt.HIGH_WICK}.get(main, "High first")
        else:
            d["first_side"] = {"Down": bt.LOW_MAIN, "Up": bt.LOW_WICK}.get(main, "Low first")
    return d


class Expected:
    def __init__(self, rows, window):
        start, end = window
        self.rows = [r for r in rows if isinstance(r.get("date"), dt.datetime) and start <= r["date"] <= end]

    def count(self, pred):
        return sum(1 for r in self.rows if pred(r))

    @staticmethod
    def ratio(a, b):
        return "–" if b == 0 else a / b

    def side_when(self, r, side):
        if r.get("main") not in ("Up", "Down"):
            return None
        main_is_high = r["main"] == "Up"
        use_high = main_is_high if side == "main" else not main_is_high
        return r.get("high_when") if use_high else r.get("low_when")

    def close_count(self, key, kind):
        pairs = {"strong": {("Up", "Top"), ("Down", "Bottom")}, "middle": {("Up", "Middle"), ("Down", "Middle")},
                 "weak": {("Up", "Bottom"), ("Down", "Top")}}[kind]
        return self.count(lambda r: (r.get("main"), r.get(key)) in pairs)

    def value(self, spec):
        kind, args = spec[0], spec[1:]
        c, ratio = self.count, self.ratio
        if kind == "releases":
            return len(self.rows)
        if kind == "n":
            key, labels, ckey, cval = args
            return c(lambda r: r.get(key) in labels and (ckey is None or cval is None or r.get(ckey) == cval))
        if kind == "cond_share":
            ckey, cval, o = args
            cond = (lambda r: True) if cval is None else (lambda r: r.get(ckey) == cval)
            return ratio(c(lambda r: cond(r) and r["outcome"] == o),
                         c(lambda r: cond(r) and r["outcome"] in bt.OUTCOMES))
        if kind == "count":
            key, label = args
            return c(lambda r: r.get(key) == label)
        if kind == "share":
            key, subset, universe = args
            return ratio(c(lambda r: r.get(key) in subset), c(lambda r: r.get(key) in universe))
        if kind == "total":
            key, labels = args
            return c(lambda r: r.get(key) in labels)
        if kind == "side":
            side, windows = args
            return c(lambda r: self.side_when(r, side) in windows)
        if kind == "side_share":
            side, upto = args
            return ratio(c(lambda r: self.side_when(r, side) in upto), c(lambda r: self.side_when(r, side) in bt.WIN_HL))
        if kind == "close":
            key, which = args
            n = sum(self.close_count(key, k) for k in ("strong", "middle", "weak"))
            return n if which is None else ratio(self.close_count(key, which), n)
        pbs = lambda cond: [r["pullback"] for r in self.rows if isnum(r.get("pullback")) and r["pullback"] >= 0
                            and (cond is None or r["outcome"] == cond)]
        if kind == "pb_avg":
            vals = pbs(args[0])
            return "–" if not vals else sum(vals) / len(vals)
        if kind == "pb_reach":
            level, cond = args
            vals = pbs(cond)
            return ratio(sum(1 for v in vals if v >= level), len(vals))
        if kind == "pb_n":
            return len(pbs(args[0]))
        if kind == "cpi_dir":
            cpi, d = args
            n = c(lambda r: r.get("cpi") == cpi and r.get("main") in ("Up", "Down"))
            return n if d is None else ratio(c(lambda r: r.get("cpi") == cpi and r.get("main") == d), n)
        if kind == "avg":
            key, d = args
            vals = [r[key] for r in self.rows if isnum(r.get(key)) and r[key] > -1e6 and (d is None or r.get("main") == d)]
            return "–" if not vals else sum(vals) / len(vals)
        if kind == "bias":
            w = c(lambda r: (r.get("bias"), r.get("main")) in {("Bullish", "Up"), ("Bearish", "Down")})
            a = c(lambda r: (r.get("bias"), r.get("main")) in {("Bullish", "Down"), ("Bearish", "Up")})
            return {"with": w, "against": a, "n": w + a, "with_share": ratio(w, w + a),
                    "against_share": ratio(a, w + a)}[args[0]]
        if kind == "setup":
            setup, what = args
            names = bt.SETUP[1:] if setup is None else [setup]
            rs = [r["r"] for r in self.rows if r.get("setup") in names and isnum(r.get("r"))]
            trades = [v for v in rs if v > -1000]
            wins = [v for v in rs if v > 0]
            if what == "trades":
                return len(trades)
            if what == "wins":
                return len(wins)
            if what == "winrate":
                return ratio(len(wins), len(trades))
            if what == "total":
                return sum(rs)
            if what == "avg":
                return "–" if not trades else (sum(rs) if setup is None else sum(trades)) / len(trades)
        raise ValueError(spec)


# ------------------------------------------------------------------ LibreOffice
def recalculate(src: Path, outdir: Path) -> Path:
    profile = outdir / "lo_profile"
    user = profile / "user"
    user.mkdir(parents=True, exist_ok=True)
    (user / "registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
        'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
        '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
        '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>\n'
        '</oor:items>\n')
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("LibreOffice (soffice) is required for verification")
    out = outdir / "calc"
    subprocess.run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--calc",
                    "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=300)
    return out / src.name


def close_enough(a, b):
    if isnum(a) and isnum(b):
        return abs(a - b) <= 1e-9 * max(1, abs(b))
    return (a if not blank(a) else "") == (b if not blank(b) else "")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--legacy", type=Path, help="xlsx export of the old tab (only A:K are used)")
    parser.add_argument("--rows", type=int, default=160)
    parser.add_argument("--keep", type=Path, help="copy the recalculated test workbook here")
    args = parser.parse_args()

    rng = random.Random(20260924)
    rows = synthetic_rows(rng, args.rows)
    if args.legacy:
        extra = legacy_rows(args.legacy)
        add_random_new_columns(extra, rng)
        rows = extra + rows
    assert len(rows) <= bt.LAST_ROW - bt.FIRST_ROW + 1, "too many test rows"

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        expect = {}
        template = tmp / "template.xlsx"
        bt.build(template, expect=expect)

        wb = load_workbook(template)
        log, stats = wb[bt.LOG], wb[bt.STATS]
        dv_before = len(log.data_validations.dataValidation)
        cf_before = sum(len(r.rules) for r in log.conditional_formatting)
        for i, row in enumerate(rows):
            r = bt.FIRST_ROW + i
            for key, value in row.items():
                if key == "day" and blank(value):
                    continue  # keep the formula, like a new row
                log[f"{bt.C[key]}{r}"].value = value
        stats["C3"].value, stats["E3"].value = WINDOW
        test = tmp / "test.xlsx"
        wb.save(test)

        calc = recalculate(test, tmp)
        if args.keep:
            shutil.copy(calc, args.keep)
        formulas = load_workbook(calc)
        values = load_workbook(calc, data_only=True)

        failures = []
        # 1. no formula anywhere evaluates to an error
        for ws in values.worksheets:
            for line in ws.iter_rows():
                for cell in line:
                    if isinstance(cell.value, str) and cell.value.startswith(ERRORS):
                        failures.append(f"{ws.title}!{cell.coordinate} = {cell.value}")

        # 2. auto columns match the Python rules
        vlog = values[bt.LOG]
        for i, row in enumerate(rows):
            r = bt.FIRST_ROW + i
            for key, want in derive(row).items():
                got = vlog[f"{bt.C[key]}{r}"].value
                if not close_enough(got, want):
                    failures.append(f"{bt.LOG}!{bt.C[key]}{r} ({key}) = {got!r}, expected {want!r}")
        blank_row = bt.FIRST_ROW + len(rows)
        for key in ("day", "move_vs_data", "outcome", "first_side"):
            if not blank(vlog[f"{bt.C[key]}{blank_row}"].value):
                failures.append(f"empty row shows {key}")

        # 3. every Stats formula matches the Python recalculation
        for row in rows:
            row.update({k: v for k, v in derive(row).items() if k != "day"})
        exp = Expected(rows, WINDOW)
        vstats = values[bt.STATS]
        for cell, spec in expect.items():
            got, want = vstats[cell].value, exp.value(spec)
            if not close_enough(got, want):
                failures.append(f"{bt.STATS}!{cell} {spec[:3]} = {got!r}, expected {want!r}")

        # 4. LibreOffice kept all dropdowns / colour rules (i.e. they parsed)
        flog = formulas[bt.LOG]
        dv_after = len(flog.data_validations.dataValidation)
        cf_after = sum(len(r.rules) for r in flog.conditional_formatting)
        if dv_after < dv_before or cf_after < cf_before:
            failures.append(f"validation/format rules lost: dv {dv_before}->{dv_after}, cf {cf_before}->{cf_after}")

        nonzero = sum(1 for spec in expect.values() if exp.value(spec) not in (0, "–"))
        print(f"rows tested: {len(rows)} · stats cells checked: {len(expect)} ({nonzero} non-zero) · "
              f"dropdowns: {dv_after} · colour rules: {cf_after}")
        if failures:
            print(f"{len(failures)} FAILURES")
            for f in failures[:60]:
                print("  " + f)
            sys.exit(1)
        print("ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
