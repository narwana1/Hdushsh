#!/usr/bin/env python3
"""Render PNG previews of the template (3 made-up example rows) with LibreOffice.

    python tools/render_previews.py --outdir screenshots
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image, ImageChops

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_template as bt  # noqa: E402


def to_png(xlsx: Path, png: Path, tmp: Path, dpi: int):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    profile = (tmp / "lo_profile").as_uri()
    subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                    "--outdir", str(tmp), str(xlsx)], check=True, capture_output=True, timeout=300)
    pdf = tmp / (xlsx.stem + ".pdf")
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-f", "1", "-l", "1", "-singlefile", str(pdf),
                    str(tmp / xlsx.stem)], check=True, timeout=300)
    img = Image.open(tmp / f"{xlsx.stem}.png").convert("RGB")
    bbox = ImageChops.difference(img, Image.new("RGB", img.size, "white")).getbbox()
    img.crop((max(0, bbox[0] - 12), max(0, bbox[1] - 12), bbox[2] + 12, bbox[3] + 12)).save(png)
    print(f"wrote {png}")


def log_view(src: Path, dst: Path, hide: tuple[str, str] | None, last_col: str):
    wb = load_workbook(src)
    for name in (bt.STATS, bt.GUIDE):
        del wb[name]
    ws = wb[bt.LOG]
    if hide:
        for idx in range(column_index_from_string(hide[0]), column_index_from_string(hide[1]) + 1):
            ws.column_dimensions[get_column_letter(idx)].hidden = True
    ws.print_area = f"A1:{last_col}{bt.FIRST_ROW + len(bt.DEMO_ROWS) - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = True
    wb.save(dst)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--outdir", type=Path, default=Path("screenshots"))
    parser.add_argument("--dpi", type=int, default=170)
    args = parser.parse_args()
    args.outdir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        demo = tmp / "demo.xlsx"
        bt.build(demo, demo=True)
        split_after = bt.C["close15"]  # left image: your columns + CPI + candle closes
        left, right = tmp / "log_left.xlsx", tmp / "log_right.xlsx"
        log_view(demo, left, None, split_after)
        log_view(demo, right, (bt.C["wick"], split_after), bt.COLUMNS[-1].letter)
        to_png(left, args.outdir / "cpi_log_part1.png", tmp, args.dpi)
        to_png(right, args.outdir / "cpi_log_part2.png", tmp, args.dpi)


if __name__ == "__main__":
    main()
