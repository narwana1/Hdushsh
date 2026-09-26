#!/usr/bin/env python3
"""Google Sheet helpers for the CPI Log tab.

    python backtest/sheet.py snapshot --out backtest/out/pre.xlsx
    python backtest/sheet.py fill  --row 82 --values backtest/out/2026-07-14.json [--link URL] [--dry-run]
    python backtest/sheet.py check --pre backtest/out/pre.xlsx --row 82 --values backtest/out/2026-07-14.json [--link URL]

Needs CPI_SHEET_ID: the id from the working sheet's URL. Ask the owner; it is not
stored in this public repo because anyone with the link can edit the sheet.

fill only writes the columns computed by rules.py (plus the chart link), never the
owner's own columns A-K or the automatic columns, and skips cells that already
have a value unless --overwrite is given. check compares a fresh export with the
snapshot taken before the edit and recomputes every Stats cell.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import sys
import tempfile
import time
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import build_template as bt  # noqa: E402
import verify_template as vt  # noqa: E402

WRITABLE = ["cpi", "close1", "close5", "close15", "high_when", "low_when", "pullback", "back_open", "at_open",
            "bias", "midnight", "liq_wick", "liq_main", "pda", "fvg", "smt", "setup", "r", "link"]
AUTO = [c.key for c in bt.COLUMNS if c.auto]
OPTIONS = {c.key: c.options for c in bt.COLUMNS if c.options}


def sheet_id():
    sid = os.environ.get("CPI_SHEET_ID", "").strip()
    if not sid:
        sys.exit("set CPI_SHEET_ID to the id in the sheet's URL (ask the owner)")
    return sid


def export(path: Path):
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id()}/export?format=xlsx&t={time.time()}"
    with urllib.request.urlopen(url, timeout=120) as r:
        path.write_bytes(r.read())
    return path


def wanted(args):
    raw = json.loads(Path(args.values).read_text()) if args.values else {}
    vals = raw.get("values", raw)
    if args.link:
        vals = {**vals, "link": args.link}
    bad = [k for k in vals if k not in WRITABLE]
    if bad:
        sys.exit(f"refusing to write {bad}: only {WRITABLE}")
    return {k: vals[k] for k in WRITABLE if vals.get(k) is not None}


def same(a, b):
    if isinstance(b, (int, float)) and not isinstance(b, bool):
        return isinstance(a, (int, float)) and not isinstance(a, bool) and abs(a - b) < 1e-9
    return a == b


# ------------------------------------------------------------------ check
def style(x):
    b = x.border
    return (x.number_format, x.font.name, x.font.sz, x.font.b, x.font.i, x.font.u,
            x.font.color.rgb if x.font.color else None, x.fill.fill_type, x.fill.fgColor.rgb,
            tuple(getattr(getattr(b, s), "style", None) if getattr(b, s) else None for s in ("left", "right", "top", "bottom")),
            x.alignment.horizontal, x.alignment.vertical, x.alignment.wrap_text)


def log_rows(ws_values):
    rows = []
    for r in range(bt.FIRST_ROW, bt.LAST_ROW + 1):
        row = {c.key: ws_values[f"{c.letter}{r}"].value for c in bt.COLUMNS}
        if any(row[k] not in (None, "") for k in row if k not in AUTO):
            rows.append((r, row))
    return rows


def check(args):
    want = wanted(args)
    targets = {f"{bt.C[k]}{args.row}": v for k, v in want.items()}
    auto_cells = {f"{bt.C[k]}{args.row}" for k in AUTO}
    post = Path(args.post) if args.post else export(Path(tempfile.mkdtemp()) / "post.xlsx")
    pre_f, post_f = load_workbook(args.pre), load_workbook(post)
    pre_v, post_v = load_workbook(args.pre, data_only=True), load_workbook(post, data_only=True)
    fails, notes = [], []

    log_v = post_v[bt.LOG]
    for ref, v in targets.items():
        if not same(log_v[ref].value, v):
            fails.append(f"{ref} = {log_v[ref].value!r}, expected {v!r}")
    for ws_pre in pre_f.worksheets:
        ws_post = post_f[ws_pre.title]
        for r in range(1, max(ws_pre.max_row, ws_post.max_row) + 1):
            for c in range(1, max(ws_pre.max_column, ws_post.max_column) + 1):
                a, b = ws_pre.cell(r, c), ws_post.cell(r, c)
                ref, mine = a.coordinate, ws_pre.title == bt.LOG
                if style(a) != style(b):
                    (notes if mine and ref in targets else fails).append(f"{ws_pre.title}!{ref} style changed")
                if mine and ref in targets:
                    continue
                if a.value != b.value:
                    fails.append(f"{ws_pre.title}!{ref}: {a.value!r} -> {b.value!r}")
                elif ws_pre.title != bt.STATS and not (mine and ref in auto_cells) and \
                        pre_v[ws_pre.title][ref].value != post_v[ws_pre.title][ref].value:
                    fails.append(f"{ws_pre.title}!{ref} value {pre_v[ws_pre.title][ref].value!r} -> "
                                 f"{post_v[ws_pre.title][ref].value!r}")
    for name in (bt.LOG, bt.STATS, bt.GUIDE):
        a, b = pre_f[name], post_f[name]
        sigs = [("data validation", lambda w: sorted((str(d.sqref), d.type, d.formula1, d.formula2)
                                                     for d in w.data_validations.dataValidation)),
                ("conditional formatting", lambda w: sorted((str(k.sqref), len(k.rules)) for k in w.conditional_formatting)),
                ("freeze", lambda w: w.freeze_panes), ("merges", lambda w: sorted(map(str, w.merged_cells.ranges)))]
        fails += [f"{name}: {label} changed" for label, f in sigs if f(a) != f(b)]
    for ws in post_v.worksheets:
        for line in ws.iter_rows():
            fails += [f"{ws.title}!{x.coordinate} = {x.value}" for x in line
                      if isinstance(x.value, str) and x.value.startswith(vt.ERRORS)]

    rows = []
    for r, row in log_rows(log_v):
        day = row["date"].strftime("%a") if isinstance(row["date"], dt.datetime) else ""
        derived = vt.derive({**row, "day": None})
        if (row["day"] or "") != day:
            fails.append(f"B{r} shows {row['day']!r}, expected {day!r}")
        for k in ("move_vs_data", "outcome", "first_side"):
            if (row[k] or "") != derived[k]:
                fails.append(f"{bt.C[k]}{r} shows {row[k]!r}, expected {derived[k]!r}")
        rows.append({**row, **{k: derived[k] for k in ("move_vs_data", "outcome", "first_side")}})
    stv = post_v[bt.STATS]
    exp = vt.Expected(rows, (stv["C3"].value, stv["E3"].value))
    spec = {}
    bt.build(Path(tempfile.mkdtemp()) / "spec.xlsx", expect=spec)
    for cell, s in spec.items():
        if not vt.close_enough(stv[cell].value, exp.value(s)):
            fails.append(f"Stats!{cell} = {stv[cell].value!r}, expected {exp.value(s)!r}")

    print(f"row {args.row}: {len(targets)} cells checked | {len(rows)} log rows | {len(spec)} Stats cells recomputed")
    for n in notes:
        print("note:", n, "(expected for the link cell: Google underlines links)")
    if fails:
        print(len(fails), "PROBLEMS")
        print("\n".join(fails[:50]))
        return 1
    print("ALL CHECKS PASSED")
    return 0


# ------------------------------------------------------------------ fill
def fill(args):
    from playwright.sync_api import sync_playwright

    want = wanted(args)
    problems, probed = [], []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_context(viewport={"width": 1600, "height": 1000}, locale="en-US").new_page()
        page.goto(f"https://docs.google.com/spreadsheets/d/{sheet_id()}/edit", wait_until="domcontentloaded", timeout=90000)
        page.wait_for_timeout(12000)
        box, bar = page.locator("#t-name-box"), page.locator("#t-formula-bar-input")

        def select(ref):
            box.click()
            box.fill(ref)
            box.press("Enter")
            page.wait_for_timeout(900)
            if box.input_value() != ref:
                raise RuntimeError(f"could not select {ref}")

        def read(ref):
            select(ref)
            return bar.inner_text().strip()

        def dialog():
            d = page.locator(".modal-dialog:visible")
            if d.count():
                text = d.first.inner_text()[:200]
                d.first.get_by_role("button", name="OK").click()
                return text
            return None

        for key, value in want.items():
            ref = f"{bt.C[key]}{args.row}"
            current = read(ref)
            if current and not args.overwrite:
                problems.append(f"{ref} ({key}) already has {current!r} - skipped")
                continue
            if key in OPTIONS:
                page.keyboard.press("Enter")  # opens the dropdown of a validated cell
                page.wait_for_timeout(1500)
                opt = page.get_by_role("option", name=str(value), exact=True)
                if opt.count() != 1:
                    page.keyboard.press("Escape")
                    problems.append(f"{ref}: option {value!r} found {opt.count()} times")
                    continue
                if args.dry_run:
                    page.keyboard.press("Escape")
                    page.wait_for_timeout(800)
                    probed.append(ref)
                    print(f"dry run: {ref} ({key}) <- {value!r} (option found)")
                    continue
                opt.last.click()
            else:
                if args.dry_run:
                    print(f"dry run: {ref} ({key}) <- {value!r}")
                    continue
                page.keyboard.type(str(value), delay=15)  # pasting a URL failed once; typing is reliable
                page.keyboard.press("Enter")
            page.wait_for_timeout(2200)
            err = dialog()
            got = read(ref)
            if key in ("pullback", "r"):
                try:  # the pullback format makes 97 show as "9700%" in the formula bar
                    num = float(got.rstrip("%")) / 100 if got.endswith("%") else float(got)
                    ok = same(num, float(value))
                except ValueError:
                    ok = False
            else:
                ok = got == str(value)
            print(f"{ref} ({key}) <- {value!r}: {'ok' if ok and not err else 'PROBLEM'}")
            if err or not ok:
                problems.append(f"{ref}: wrote {value!r}, cell shows {got!r}" + (f", dialog: {err}" if err else ""))
        for ref in probed:  # opening a dropdown in a dry run must not write anything
            after = read(ref)
            if after:
                problems.append(f"dry run left {after!r} in {ref}")
        browser.close()
    for pr in problems:
        print("PROBLEM:", pr)
    return 1 if problems else 0


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("snapshot", help="download the sheet as xlsx (do this before every edit)")
    s.add_argument("--out", required=True, type=Path)
    for name in ("fill", "check"):
        s = sub.add_parser(name)
        s.add_argument("--row", required=True, type=int, help="CPI Log row number, e.g. 82")
        s.add_argument("--values", help="rules.py output json")
        s.add_argument("--link", help="chart link for column AG")
        if name == "fill":
            s.add_argument("--dry-run", action="store_true")
            s.add_argument("--overwrite", action="store_true")
        else:
            s.add_argument("--pre", required=True, help="snapshot taken before the edit")
            s.add_argument("--post", help="use this export instead of downloading a fresh one")
    args = ap.parse_args(argv)
    if args.cmd == "snapshot":
        args.out.parent.mkdir(parents=True, exist_ok=True)
        print("wrote", export(args.out))
        return 0
    if not bt.FIRST_ROW <= args.row <= bt.LAST_ROW:
        sys.exit(f"row must be {bt.FIRST_ROW}-{bt.LAST_ROW}")
    return fill(args) if args.cmd == "fill" else check(args)


if __name__ == "__main__":
    sys.exit(main())
